"""File management tools for the ReAct agent: read_file and write_code_file with smart path resolution and ambiguity handling."""

import os
from pathlib import Path
from typing import List, Optional
from langchain_core.tools import tool

# Supported text extensions
TEXT_EXTENSIONS = {
    ".txt", ".md", ".py", ".json", ".csv", ".xml", ".html", ".css",
    ".js", ".ts", ".yml", ".yaml", ".ini", ".cfg", ".log", ".env"
}


def find_matching_paths(raw_path: str) -> List[Path]:
    """Finds all matching file paths across workspace, Desktop, Downloads, and Documents.
    Returns exact match if found, or all fuzzy candidates matching the query.
    """
    clean_str = os.path.expanduser(os.path.expandvars(raw_path.strip().strip("'\"")))
    direct_path = Path(clean_str)

    # 1. Exact direct path check
    if direct_path.is_file():
        return [direct_path.resolve()]

    home = Path.home()
    search_dirs = [
        Path.cwd(),
        home / "Desktop",
        home / "Downloads",
        home / "Documents",
        home,
    ]

    # Check for direct relative path in search roots
    exact_matches = []
    for root in search_dirs:
        candidate = (root / clean_str).resolve()
        if candidate.is_file() and candidate not in exact_matches:
            exact_matches.append(candidate)

    if exact_matches:
        return exact_matches

    # 2. Fuzzy / partial match across search roots
    stem_query = Path(clean_str).stem.lower()
    matches: List[Path] = []
    
    for root in search_dirs:
        if not root.exists() or not root.is_dir():
            continue
        try:
            for item in root.iterdir():
                if item.is_file() and item not in matches:
                    item_stem = item.stem.lower()
                    item_name = item.name.lower()
                    if stem_query == item_stem or stem_query in item_name or stem_query in item_stem:
                        matches.append(item.resolve())
        except (PermissionError, OSError):
            continue

    return matches


@tool
def read_file(path: str) -> str:
    """Reads and returns the contents of a local file (.txt, .md, .py, .pdf, .docx, .xlsx).
    Smartly searches the workspace, Desktop, Downloads, and Documents if a relative path or keyword is given.
    If multiple files match the keyword, returns a list of matching files for user selection.
    
    Args:
        path: Relative, absolute, or partial path/keyword to the file to read (e.g. 'notes.txt', 'resume').
        
    Returns:
        The text content of the file, a list of matching options if ambiguous, or a descriptive error message.
    """
    try:
        clean_str = os.path.expanduser(os.path.expandvars(path.strip().strip("'\"")))
        direct_path = Path(clean_str)
        if direct_path.exists() and direct_path.is_dir():
            return f"Error: Path '{path}' is a directory, not a file."

        matches = find_matching_paths(path)
        
        if not matches:
            return f"Error: File not found at path '{path}' (searched workspace, Desktop, Downloads, Documents)."
            
        if len(matches) > 1:
            options = "\n".join([f"{i+1}. {m.name} (Path: {m})" for i, m in enumerate(matches[:10])])
            return (
                f"Multiple matching files found for '{path}':\n{options}\n\n"
                f"Please ask the user which specific file they would like to read."
            )

        file_path = matches[0]
        
        if file_path.is_dir():
            return f"Error: Path '{path}' is a directory, not a file."
            
        suffix = file_path.suffix.lower()
        
        # 1. Plain text / code / markdown files
        if suffix in TEXT_EXTENSIONS or suffix == "":
            try:
                return file_path.read_text(encoding="utf-8")
            except UnicodeDecodeError:
                return file_path.read_text(encoding="latin-1", errors="replace")
                
        # 2. PDF documents
        if suffix == ".pdf":
            import pdfplumber
            extracted_text = []
            with pdfplumber.open(file_path) as pdf:
                for i, page in enumerate(pdf.pages):
                    page_text = page.extract_text() or ""
                    extracted_text.append(f"--- Page {i + 1} ---\n{page_text}")
            return "\n\n".join(extracted_text) if extracted_text else "Note: PDF was empty or contained only images."

        # 3. DOCX documents
        if suffix == ".docx":
            import docx
            doc = docx.Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            return "\n".join(paragraphs) if paragraphs else "Note: DOCX file was empty."

        # 4. Excel spreadsheets
        if suffix in [".xlsx", ".xls"]:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            output = []
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                output.append(f"=== Sheet: {sheet_name} ===")
                for row in sheet.iter_rows(values_only=True):
                    row_vals = [str(c) if c is not None else "" for c in row]
                    if any(row_vals):
                        output.append("\t".join(row_vals))
            return "\n".join(output) if output else "Note: Spreadsheet was empty."

        return f"Error: Unsupported file format '{suffix}' for reading."

    except Exception as e:
        return f"Error reading file '{path}': {str(e)}"


@tool
def write_code_file(filepath: str, code_content: str) -> str:
    """Generates or overwrites a code/text file at the specified path.
    
    Args:
        filepath: Relative or absolute path where the file should be saved.
        code_content: The full content/source code to write to the file.
        
    Returns:
        Success confirmation message with path and byte size, or an error message.
    """
    try:
        clean_str = os.path.expanduser(os.path.expandvars(filepath.strip().strip("'\"")))
        target_path = Path(clean_str).resolve()
        
        # Ensure parent directories exist
        target_path.parent.mkdir(parents=True, exist_ok=True)
        
        target_path.write_text(code_content, encoding="utf-8")
        byte_size = target_path.stat().st_size
        return f"Success: File '{filepath}' written successfully ({byte_size} bytes)."
        
    except Exception as e:
        return f"Error writing file '{filepath}': {str(e)}"
